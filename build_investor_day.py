import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyArrowPatch
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ── COLORS ───────────────────────────────────────────────────────────────────
NAVY   = "#1A3A5C"
TEAL   = "#2A7F8C"
RED    = "#C0392B"
ORANGE = "#E67E22"
GREY   = "#7F8C8D"
GREEN  = "#27AE60"
LGREY  = "#ECF0F1"

# ── 2×2 POSITIONING MATRIX ───────────────────────────────────────────────────
# X-axis: PM-centric (0) → Agentic (10)
# Y-axis: Bundled pricing (0) → Premium/consumption pricing (10)

companies = {
    "Asana":      {"x": 7.5, "y": 3.5, "color": ORANGE, "label_dx": 0.15, "label_dy": 0.3},
    "Monday":     {"x": 5.5, "y": 3.0, "color": RED,    "label_dx": 0.15, "label_dy": -0.5},
    "Smartsheet": {"x": 3.0, "y": 5.5, "color": GREY,   "label_dx": 0.15, "label_dy": 0.3},
    "Atlassian":  {"x": 8.5, "y": 8.0, "color": GREEN,  "label_dx": 0.15, "label_dy": 0.3},
    "Meridian\n(Recommended)": {"x": 7.8, "y": 7.5, "color": NAVY,  "label_dx": -2.3, "label_dy": 0.35},
}

fig, ax = plt.subplots(figsize=(10, 8))
fig.patch.set_facecolor("#F9F9F9")
ax.set_facecolor("#F9F9F9")

# Quadrant shading
ax.axhspan(5, 10, xmin=0.5, xmax=1.0, alpha=0.06, color=NAVY)   # top-right: target zone
ax.axvline(5, color="#CCCCCC", lw=1.2, ls="--")
ax.axhline(5, color="#CCCCCC", lw=1.2, ls="--")

# Quadrant labels
ax.text(2.5, 9.5, "Governance-first\n(cautious)", ha="center", fontsize=9,
        color="#AAAAAA", style="italic")
ax.text(7.5, 9.5, "Agentic + Premium\n★ Target zone", ha="center", fontsize=9,
        color=NAVY, style="italic", fontweight="bold")
ax.text(2.5, 0.5, "PM-centric\nbundled", ha="center", fontsize=9, color="#AAAAAA", style="italic")
ax.text(7.5, 0.5, "Agentic\nbundled", ha="center", fontsize=9, color="#AAAAAA", style="italic")

# Plot companies
for name, d in companies.items():
    size = 320 if "Meridian" in name else 200
    ax.scatter(d["x"], d["y"], s=size, color=d["color"],
               zorder=5, edgecolors="white", linewidths=1.5)
    ax.text(d["x"] + d["label_dx"], d["y"] + d["label_dy"],
            name, fontsize=10, fontweight="bold" if "Meridian" in name else "normal",
            color=d["color"], va="center")

# Arrow showing Meridian's direction of travel
ax.annotate("", xy=(7.8, 7.5), xytext=(6.0, 5.8),
            arrowprops=dict(arrowstyle="->", color=NAVY, lw=1.8, ls="dashed"))
ax.text(5.8, 5.5, "Meridian's\ntrajectory", fontsize=8, color=NAVY, style="italic")

ax.set_xlim(0, 10)
ax.set_ylim(0, 10)
ax.set_xlabel("PM-centric  ◄──────────────────►  Agentic Platform", fontsize=11, labelpad=10)
ax.set_ylabel("Bundled pricing  ◄──────────────────►  Premium / Consumption pricing",
              fontsize=11, labelpad=10)
ax.set_title("Competitive Positioning Matrix — Work Management + AI (Feb 2026)",
             fontsize=13, fontweight="bold", color=NAVY, pad=14)
ax.set_xticks([])
ax.set_yticks([])
ax.spines["top"].set_visible(False)
ax.spines["right"].set_visible(False)

plt.tight_layout()
matrix_path = "/home/user/ClaudeCodeBerkeley/competitive_matrix.png"
plt.savefig(matrix_path, dpi=160, bbox_inches="tight", facecolor=fig.get_facecolor())
plt.close()
print(f"Matrix saved: {matrix_path}")


# ── WORD DOCUMENT ─────────────────────────────────────────────────────────────

doc = Document()
for section in doc.sections:
    section.top_margin    = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin   = Inches(1.0)
    section.right_margin  = Inches(1.0)

def hp(doc, text, level=1, color=(26,58,92)):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(*color)
    return p

def para(doc, text, size=11, bold=False, italic=False, color=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p

def bullet(doc, text, bold_prefix=None, size=11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.font.name = "Calibri"; r1.font.size = Pt(size); r1.font.bold = True
        r2 = p.add_run(text)
        r2.font.name = "Calibri"; r2.font.size = Pt(size)
    else:
        r = p.add_run(text)
        r.font.name = "Calibri"; r.font.size = Pt(size)

# ── HEADER ───────────────────────────────────────────────────────────────────
h = doc.add_heading("", 0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run("MERIDIAN TECHNOLOGIES")
r.font.name = "Calibri"; r.font.size = Pt(20); r.font.bold = True
r.font.color.rgb = RGBColor(26,58,92)

s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
s.paragraph_format.space_after = Pt(2)
r = s.add_run("Investor Day Positioning Memo  |  March 11, 2026")
r.font.name="Calibri"; r.font.size=Pt(12); r.font.color.rgb=RGBColor(42,127,140)

m = doc.add_paragraph(); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
m.paragraph_format.space_after = Pt(16)
r = m.add_run("Catherine Park, President & CEO  |  CONFIDENTIAL — Pre-release draft")
r.font.name="Calibri"; r.font.size=Pt(10); r.font.color.rgb=RGBColor(127,140,141)

# ── EXEC SUMMARY ─────────────────────────────────────────────────────────────
hp(doc, "Executive Summary", 2)

para(doc,
    "Meridian will declare itself the enterprise-governed agentic work platform — "
    "the company that combines the Helio agent framework with the deepest enterprise "
    "governance stack in the category. Project management is our wedge and our proven "
    "foundation; agents running on trusted, auditable, regulation-ready infrastructure "
    "are our three-year strategy. We are not building AI features on top of a PM tool. "
    "We are building the platform on which regulated enterprises can run agents safely "
    "— and no competitor owns that position today.",
    size=11, space_after=12)

# ── POSITION ─────────────────────────────────────────────────────────────────
hp(doc, "Our Position (One Sentence for the Market)", 2)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
r = p.add_run(
    '"Meridian is the agentic work platform that regulated enterprises trust — '
    'built on the governance infrastructure that banks, pharma, and federal agencies '
    'already rely on, and powered by the Helio agent framework."'
)
r.font.name="Calibri"; r.font.size=Pt(12); r.font.bold=True; r.font.italic=True
r.font.color.rgb=RGBColor(26,58,92)

# ── WHY ──────────────────────────────────────────────────────────────────────
hp(doc, "Why This Position — Three Reasons", 2)

bullet(doc,
    "The competitive white space is here. Asana and Monday are agentic but have no governance. "
    "Smartsheet has governance but explicitly rejects agentic. Atlassian is agentic + governance "
    "but only for software/IT teams. No one owns 'agentic + regulated-industry governance.' "
    "We do, today.",
    bold_prefix="1.  The white space is ours. ")

bullet(doc,
    "18 of 22 enterprise advisory board sessions in the last 90 days asked for agent auditability, "
    "role-based permissions, model selection, and data residency. These customers are not asking "
    "for better PM features. They are asking for a governed agent platform. We should name what "
    "we already are.",
    bold_prefix="2.  Our customers are already asking for it. ")

bullet(doc,
    "We have $506M liquidity, $55M AI R&D budgeted, and the Helio team — 28 engineers "
    "from leading AI labs with a working agent framework. Option A (PM-with-AI) would "
    "likely cost us the Helio team at their 2026 compensation cliff. Option B keeps them "
    "and uses their full capacity.",
    bold_prefix="3.  We have the assets to execute it. ")

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ── MATRIX CHART ─────────────────────────────────────────────────────────────
hp(doc, "Competitive Landscape", 2)

para(doc,
    "The 2×2 below maps all four direct competitors on two dimensions: how agentic their "
    "public positioning is (X-axis) and how premium their AI pricing model is (Y-axis). "
    "The upper-right quadrant — agentic positioning with premium/consumption pricing — "
    "is the target zone. Only Atlassian and Meridian are positioned there; Atlassian's "
    "footprint is in software/IT, not regulated general enterprise.",
    size=11, space_after=6)

doc.add_picture(matrix_path, width=Inches(5.8))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_after = Pt(10)
r = cap.add_run("Figure 1 — Competitive positioning matrix. Meridian's recommended trajectory shown by dashed arrow.")
r.font.name="Calibri"; r.font.size=Pt(9); r.font.color.rgb=RGBColor(127,140,141)

# ── COMPETITOR TABLE ─────────────────────────────────────────────────────────
hp(doc, "Competitor AI Positioning — Summary Table", 2)

table = doc.add_table(rows=6, cols=5)
table.style = "Table Grid"

headers = ["Company", "AI Positioning", "Pricing", "Flagship Announcement", "vs. Meridian"]
rows_data = [
    ["Asana",      "Work mgmt + AI built in; 'agent OS' language",
     "Bundled Advanced+",  "AI Studio + Smart Workflows bundled, Nov 2025",
     "Agentic, no governance moat"],
    ["Monday.com", "Work OS 'supercharged with AI'; breadth play",
     "Bundled Pro+",       "monday AI Agents GA, Jan 2026",
     "Price aggressor, no regulated-industry depth"],
    ["Smartsheet", "Enterprise trust + AI; avoids 'agentic'",
     "AI bundled + Compliance Pack add-on",
     "AI Compliance Pack, Dec 2025",
     "Closest to Option A; decelerating growth"],
    ["Atlassian",  "Agentic enterprise platform; Rovo Studio",
     "Per-seat + consumption",
     "Rovo Studio GA, Jan 2026",
     "Closest peer to Option B; software/IT focus only"],
    ["Meridian\n(target)", "Enterprise-governed agentic work platform",
     "Per-seat + consumption H2 2026",
     "Copilot GA + Helio integration; governance suite Q1 2026",
     "Owns regulated-industry agentic governance"],
]

# Header row
hrow = table.rows[0]
hrow.cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)
for i, h_text in enumerate(headers):
    cell = hrow.cells[i]
    cell.paragraphs[0].clear()
    p = cell.paragraphs[0]
    r = p.add_run(h_text)
    r.font.bold = True; r.font.size = Pt(9); r.font.name = "Calibri"
    r.font.color.rgb = RGBColor(255,255,255)
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1A3A5C')
    tcPr.append(shd)

for row_idx, row_data in enumerate(rows_data):
    row = table.rows[row_idx + 1]
    is_meridian = row_idx == 4
    for col_idx, cell_text in enumerate(row_data):
        cell = row.cells[col_idx]
        cell.paragraphs[0].clear()
        p = cell.paragraphs[0]
        r = p.add_run(cell_text)
        r.font.size = Pt(8.5); r.font.name = "Calibri"
        if is_meridian:
            r.font.bold = True
            r.font.color.rgb = RGBColor(26,58,92)
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'E8F4FD')
            tcPr.append(shd)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ── THREE RISKS ───────────────────────────────────────────────────────────────
hp(doc, "Three Biggest Risks", 2)

bullet(doc,
    "Atlassian has more R&D, bigger install base, and 6–12 months head start in agentic. "
    "If they invest in regulated-industry certifications (FedRAMP High, GxP), our moat narrows. "
    "Mitigation: accelerate GxP in 2026, add a second regulated vertical (federal or insurance).",
    bold_prefix="Risk 1 — Atlassian out-executes us in our own white space. ")

bullet(doc,
    "Enterprise consumption pricing ramps slower than projected. Regulated-industry procurement "
    "cycles are 9–14 months. If per-consumption revenue is not measurable by Q4 2026, "
    "we will face a credibility gap with investors. Mitigation: announce pricing model at Investor "
    "Day with specific 2026 milestones we are willing to be held to.",
    bold_prefix="Risk 2 — Consumption pricing ramps slowly in regulated enterprise. ")

bullet(doc,
    "PM motion softens while agentic ramps. Mid-market NRR is already at 102%. If Copilot "
    "bundling for mid-market is not decided by Q2, renewals will continue to erode. "
    "Mitigation: separate the pricing decision for enterprise (hold at $40/seat) vs. "
    "mid-market (consider bundled or heavily discounted).",
    bold_prefix="Risk 3 — We concede mid-market while agentic ramp takes hold. ")

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ── THREE COMMITMENTS ─────────────────────────────────────────────────────────
hp(doc, "Three Commitments to Investors", 2)

bullet(doc,
    "Copilot attach rate on enterprise renewals stays above 40% through 2026. "
    "We will report this metric publicly starting Q1 2026.",
    bold_prefix="1.  AI adoption is measurable and we will be held to it. ")

bullet(doc,
    "Consumption pricing goes live in H2 2026. We will share the pricing architecture "
    "at Investor Day and report consumption ARR as a separate line starting Q3 2026.",
    bold_prefix="2.  Consumption pricing is not a roadmap item — it is a 2026 commitment. ")

bullet(doc,
    "GxP certification ships in Q4 2026. One additional regulated vertical (federal or "
    "insurance) enters the pipeline by Q2 2026. We are not a two-vertical company.",
    bold_prefix="3.  We expand the regulated-industry footprint — not just deepen it. ")

# ── FOOTER ───────────────────────────────────────────────────────────────────
doc.add_paragraph()
fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("CONFIDENTIAL — For Internal Use and Board Review Only  |  Meridian Technologies, Inc.  |  NASDAQ: MRDN")
r.font.name="Calibri"; r.font.size=Pt(9); r.font.color.rgb=RGBColor(127,140,141)

docx_path = "/home/user/ClaudeCodeBerkeley/investor_day_positioning_memo.docx"
doc.save(docx_path)
print(f"Word document saved: {docx_path}")
